import os
import sys
import re
import copy
import json
import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from typing import Dict, Any, List, Optional, Tuple
from order_parser import clean_product_text

# Base application directory
if getattr(sys, 'frozen', False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))

DEFAULT_EXCEL_FILENAME = "Bosta Upload Template AR V3.5.xlsx"
CLEAN_TEMPLATE_FILENAME = "Bosta_Template_Clean.xlsx"
META_FILENAME = "orders_meta.json"

class ExcelManager:
    def __init__(self, file_path: Optional[str] = None):
        if not file_path:
            self.file_path = os.path.join(APP_DIR, DEFAULT_EXCEL_FILENAME)
        else:
            self.file_path = os.path.abspath(file_path)

        self.clean_template_path = os.path.join(APP_DIR, CLEAN_TEMPLATE_FILENAME)
        self.meta_path = os.path.join(APP_DIR, META_FILENAME)

        # Ensure file exists
        self._ensure_file_exists()

        self.orders: List[Dict[str, Any]] = []
        self.phone_index: Dict[str, List[Dict[str, Any]]] = {}
        self.name_index: Dict[str, List[Dict[str, Any]]] = {}
        self.meta_data: Dict[str, Any] = {}
        self.max_ord_number = 0
        self.last_filled_row = 2

        self.load_metadata()
        self.reload()

    def _ensure_file_exists(self):
        """Creates working file from clean template if it does not exist."""
        if not os.path.exists(self.file_path):
            if os.path.exists(self.clean_template_path):
                import shutil
                shutil.copyfile(self.clean_template_path, self.file_path)

    def load_metadata(self):
        if os.path.exists(self.meta_path):
            try:
                with open(self.meta_path, 'r', encoding='utf-8') as f:
                    self.meta_data = json.load(f)
            except Exception:
                self.meta_data = {}
        else:
            self.meta_data = {}

    def save_metadata(self):
        try:
            with open(self.meta_path, 'w', encoding='utf-8') as f:
                json.dump(self.meta_data, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def reload(self):
        """Loads all existing orders from the template sheet 'Add orders here'."""
        self.orders.clear()
        self.phone_index.clear()
        self.name_index.clear()
        self.max_ord_number = 0
        self.last_filled_row = 2

        if not os.path.exists(self.file_path):
            return

        wb = openpyxl.load_workbook(self.file_path, data_only=True)
        if "Add orders here" not in wb.sheetnames:
            wb.close()
            return

        ws = wb["Add orders here"]
        
        for r in range(3, ws.max_row + 1):
            name = ws.cell(r, 1).value
            phone = ws.cell(r, 2).value
            second_phone = ws.cell(r, 3).value
            city = ws.cell(r, 4).value
            area = ws.cell(r, 5).value
            address = ws.cell(r, 6).value
            landmark = ws.cell(r, 7).value
            notes = ws.cell(r, 10).value
            cod = ws.cell(r, 12).value
            items = ws.cell(r, 13).value
            desc = ws.cell(r, 14).value
            ref = ws.cell(r, 15).value
            val = ws.cell(r, 16).value

            if name or phone or address:
                self.last_filled_row = r
                
                norm_p1 = self._normalize_phone_str(phone)
                norm_p2 = self._normalize_phone_str(second_phone)

                # Lookup date from metadata or default to today
                row_key = str(r)
                meta_item = self.meta_data.get(row_key, {})
                order_date = meta_item.get("date", datetime.date.today().strftime("%Y-%m-%d"))

                order_obj = {
                    "row": r,
                    "name": str(name).strip() if name else "",
                    "phone": norm_p1,
                    "raw_phone": phone,
                    "second_phone": norm_p2,
                    "raw_second_phone": second_phone,
                    "city": str(city).strip() if city else "",
                    "area": str(area).strip() if area else "",
                    "address": str(address).strip() if address else "",
                    "landmark": str(landmark).strip() if landmark else "",
                    "notes": str(notes).strip() if notes else "",
                    "cod": cod,
                    "items": items,
                    "desc": str(desc).strip() if desc else "",
                    "ref": str(ref).strip() if ref else "",
                    "val": val,
                    "date": order_date
                }

                self.orders.append(order_obj)

                # Index by phone
                if norm_p1:
                    self.phone_index.setdefault(norm_p1, []).append(order_obj)
                if norm_p2:
                    self.phone_index.setdefault(norm_p2, []).append(order_obj)

                # Index by name
                if name:
                    cleaned_name = str(name).strip().lower()
                    self.name_index.setdefault(cleaned_name, []).append(order_obj)

                # Check ORD-XXX numbering
                if notes:
                    m = re.search(r'ORD-(\d+)', str(notes), re.IGNORECASE)
                    if m:
                        num = int(m.group(1))
                        if num > self.max_ord_number:
                            self.max_ord_number = num

        wb.close()

    def _normalize_phone_str(self, p: Any) -> str:
        if p is None:
            return ""
        s = re.sub(r'[^\d]', '', str(p))
        if len(s) == 10 and s.startswith(('10', '11', '12', '15')):
            s = '0' + s
        return s

    def get_next_order_code(self) -> str:
        next_num = self.max_ord_number + 1
        return f"ORD-{next_num}"

    def check_duplicates_and_validations(self, data: Dict[str, Any]) -> Dict[str, Any]:
        issues = {
            "is_valid": True,
            "has_duplicates": False,
            "errors": [],
            "warnings": [],
            "duplicates": []
        }

        # 1. Validate Phone
        p1 = self._normalize_phone_str(data.get("phone", ""))
        if not p1:
            issues["errors"].append("رقم الهاتف مطلوب!")
            issues["is_valid"] = False
        else:
            if len(p1) != 11:
                issues["errors"].append(f"رقم الهاتف ناقص أو زائد! عدد الأرقام ({len(p1)}) بدلاً من 11 رقم.")
                issues["is_valid"] = False
            elif not p1.startswith(('010', '011', '012', '015')):
                issues["errors"].append(f"رقم الهاتف غير صالح ({p1})! يجب أن يبدأ بـ 010 أو 011 أو 012 أو 015.")
                issues["is_valid"] = False
            else:
                if p1 in self.phone_index:
                    existing = self.phone_index[p1]
                    issues["has_duplicates"] = True
                    for ex in existing:
                        issues["duplicates"].append({
                            "field": "phone",
                            "message": f"رقم الهاتف ({p1}) مسجل مسبقاً في السطر {ex['row']} باسم '{ex['name']}'",
                            "existing": ex
                        })

        # 2. Check Second Phone
        p2 = self._normalize_phone_str(data.get("second_phone", ""))
        if p2:
            if p2 in self.phone_index:
                existing = self.phone_index[p2]
                issues["has_duplicates"] = True
                for ex in existing:
                    issues["duplicates"].append({
                        "field": "second_phone",
                        "message": f"الهاتف الثاني ({p2}) مسجل مسبقاً في السطر {ex['row']} باسم '{ex['name']}'",
                        "existing": ex
                    })

        # 3. Validate Name
        name = str(data.get("name", "")).strip()
        if not name:
            issues["errors"].append("اسم العميل مطلوب!")
            issues["is_valid"] = False
        else:
            cleaned_name = name.lower()
            if cleaned_name in self.name_index:
                existing = self.name_index[cleaned_name]
                issues["warnings"].append(f"تنبيه: يوجد عميل سابق بنفس الاسم تماماً ({name}) في السطر {existing[0]['row']}.")

        # 4. Validate City
        city = str(data.get("city", "")).strip()
        if not city:
            issues["errors"].append("المحافظة / المدينة مطلوبة!")
            issues["is_valid"] = False

        # 5. Validate Address
        addr = str(data.get("address", "")).strip()
        if not addr:
            issues["errors"].append("العنوان مطلوب!")
            issues["is_valid"] = False
        elif len(addr) < 5:
            issues["warnings"].append("العنوان قصير جداً، يرجى التأكد من كتابة تفاصيل الشارع والعمارة.")

        # 6. Validate COD
        cod = str(data.get("cod_amount", "")).strip()
        if not cod:
            issues["warnings"].append("قيمة التحصيل فارغة! (سيتم تسجيلها كطلب بدون تحصيل نقدي).")
        else:
            try:
                num = float(cod)
                if num < 0:
                    issues["errors"].append("قيمة التحصيل لا يمكن أن تكون سالبة!")
                    issues["is_valid"] = False
            except ValueError:
                issues["errors"].append(f"قيمة التحصيل غير صالحة ({cod}) - يجب أن تكون رقماً!")
                issues["is_valid"] = False

        return issues

    def append_order(self, order_data: Dict[str, Any]) -> Tuple[bool, str, int]:
        if not os.path.exists(self.file_path):
            return False, f"ملف الإكسيل غير موجود: {self.file_path}", -1

        try:
            wb = openpyxl.load_workbook(self.file_path)
            if "Add orders here" not in wb.sheetnames:
                wb.close()
                return False, "ورقة العمل 'Add orders here' غير موجودة في الملف!", -1

            ws = wb["Add orders here"]

            target_row = self.last_filled_row + 1
            while any(ws.cell(target_row, c).value is not None for c in range(1, 15)):
                target_row += 1

            name = str(order_data.get("name", "")).strip()
            raw_phone = self._normalize_phone_str(order_data.get("phone", ""))
            phone_val = int(raw_phone) if raw_phone.isdigit() else raw_phone

            raw_p2 = self._normalize_phone_str(order_data.get("second_phone", ""))
            p2_val = int(raw_p2) if (raw_p2 and raw_p2.isdigit()) else (raw_p2 or None)

            city = str(order_data.get("city", "")).strip()
            area = str(order_data.get("area", "")).strip()
            address = str(order_data.get("address", "")).strip()
            landmark = str(order_data.get("landmark", "")).strip() or None
            
            raw_notes = str(order_data.get("notes", "")).strip()
            notes = clean_product_text(raw_notes) if raw_notes else None
            if not notes:
                notes = None

            try:
                cod_val = float(order_data.get("cod_amount", 0)) if order_data.get("cod_amount") else None
                if cod_val is not None and cod_val.is_integer():
                    cod_val = int(cod_val)
            except ValueError:
                cod_val = None

            try:
                items_val = int(order_data.get("items_count", 1)) if order_data.get("items_count") else 1
            except ValueError:
                items_val = 1

            raw_desc = str(order_data.get("description", "")).strip()
            desc = clean_product_text(raw_desc) or None
            val_amount = cod_val if cod_val is not None else None

            # Date
            order_date = order_data.get("date", datetime.date.today().strftime("%Y-%m-%d"))

            # Column mapping according to Bosta Template:
            # Col 10 (J): Delivery notes only (NO code)
            # Col 12 (L): Cash collection / COD amount (Price)
            # Col 14 (N): Shipment description (Products only)
            # Col 15 (O): Order reference (None - no code anywhere)
            col_map = {
                1: name,
                2: phone_val,
                3: p2_val,
                4: city,
                5: area,
                6: address,
                7: landmark,
                10: notes,
                12: cod_val,
                13: items_val,
                14: desc,
                15: None,
                16: val_amount
            }

            font_style = Font(name="Arial", size=12)
            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )

            for col_idx in range(1, 27):
                cell = ws.cell(target_row, col_idx)
                if col_idx in col_map:
                    cell.value = col_map[col_idx]
                cell.font = font_style
                cell.alignment = align_center
                cell.border = thin_border

            ws.row_dimensions[target_row].height = 28

            wb.save(self.file_path)
            wb.close()

            # Record in metadata
            self.meta_data[str(target_row)] = {
                "date": order_date,
                "notes": notes or "",
                "name": name,
                "desc": desc or "",
                "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }
            self.save_metadata()

            self.reload()

            return True, f"تمت إضافة الطلب بنجاح في السطر {target_row}!", target_row

        except PermissionError:
            return False, "فشل الحفظ: ملف الإكسيل مفتوح في برنامج آخر. يرجى إغلاقه والمحاولة ثانية.", -1
        except Exception as e:
            return False, f"حدث خطأ أثناء حفظ الطلب في الإكسيل: {str(e)}", -1

    def delete_order(self, row_number: int) -> Tuple[bool, str]:
        """Deletes a specific order row and shifts remaining rows up."""
        if row_number < 3:
            return False, "لا يمكن حذف صفوف العناوين الرئيسية!"

        try:
            wb = openpyxl.load_workbook(self.file_path)
            ws = wb["Add orders here"]

            ws.delete_rows(row_number, 1)

            wb.save(self.file_path)
            wb.close()

            # Re-index metadata
            new_meta = {}
            for r_str, val in self.meta_data.items():
                r_int = int(r_str)
                if r_int < row_number:
                    new_meta[r_str] = val
                elif r_int > row_number:
                    new_meta[str(r_int - 1)] = val
            self.meta_data = new_meta
            self.save_metadata()

            self.reload()
            return True, f"تم حذف السطر {row_number} بنجاح!"
        except PermissionError:
            return False, "فشل الحذف: ملف الإكسيل مفتوح في برنامج آخر. يرجى إغلاقه أولاً."
        except Exception as e:
            return False, f"خطأ أثناء حذف السطر: {str(e)}"

    def clear_all_orders(self) -> Tuple[bool, str]:
        """Clears all orders from the sheet while keeping template intact."""
        try:
            wb = openpyxl.load_workbook(self.file_path)
            ws = wb["Add orders here"]

            for r in range(3, ws.max_row + 1):
                for c in range(1, 30):
                    ws.cell(r, c).value = None

            wb.save(self.file_path)
            wb.close()

            self.meta_data.clear()
            self.save_metadata()

            self.reload()
            return True, "تم تفريغ كافة البيانات والبدء بشيت جديد بنجاح!"
        except PermissionError:
            return False, "فشل التفريغ: ملف الإكسيل مفتوح في برنامج آخر. يرجى إغلاقه أولاً."
        except Exception as e:
            return False, f"خطأ أثناء تفريغ الشيت: {str(e)}"

    def get_orders_by_date(self, target_date: str) -> List[Dict[str, Any]]:
        """Returns all orders matching a specific YYYY-MM-DD date."""
        return [o for o in self.orders if o.get("date") == target_date]

    def export_orders_by_date(self, target_date: str, output_path: str) -> Tuple[bool, str, int]:
        """
        Exports orders of a specific day to a brand new Bosta template.
        Returns (success, message, exported_count).
        """
        matching_orders = self.get_orders_by_date(target_date)
        if not matching_orders:
            return False, f"لا توجد طلبات مسجلة بتاريخ {target_date}!", 0

        try:
            # Use clean template as base
            src_template = self.clean_template_path if os.path.exists(self.clean_template_path) else self.file_path
            wb = openpyxl.load_workbook(src_template)
            ws = wb["Add orders here"]

            # Clear any content in rows 3+
            for r in range(3, ws.max_row + 1):
                for c in range(1, 30):
                    ws.cell(r, c).value = None

            font_style = Font(name="Arial", size=12)
            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )

            # Write matching orders starting at row 3
            current_r = 3
            for ord_obj in matching_orders:
                raw_phone = self._normalize_phone_str(ord_obj.get("phone", ""))
                phone_val = int(raw_phone) if raw_phone.isdigit() else raw_phone

                raw_p2 = self._normalize_phone_str(ord_obj.get("second_phone", ""))
                p2_val = int(raw_p2) if (raw_p2 and raw_p2.isdigit()) else (raw_p2 or None)

                cod_val = ord_obj.get("cod")
                items_val = ord_obj.get("items") or 1

                export_desc = clean_product_text(ord_obj.get("desc") or "") or None
                raw_notes = ord_obj.get("notes") or ""
                export_notes = clean_product_text(raw_notes) or None

                col_map = {
                    1: ord_obj.get("name"),
                    2: phone_val,
                    3: p2_val,
                    4: ord_obj.get("city"),
                    5: ord_obj.get("area"),
                    6: ord_obj.get("address"),
                    7: ord_obj.get("landmark") or None,
                    10: export_notes,
                    12: cod_val,
                    13: items_val,
                    14: export_desc,
                    15: None,
                    16: cod_val
                }

                for c_idx in range(1, 27):
                    cell = ws.cell(current_r, c_idx)
                    if c_idx in col_map:
                        cell.value = col_map[c_idx]
                    cell.font = font_style
                    cell.alignment = align_center
                    cell.border = thin_border

                ws.row_dimensions[current_r].height = 28
                current_r += 1

            # Ensure output directory exists
            out_dir = os.path.dirname(os.path.abspath(output_path))
            os.makedirs(out_dir, exist_ok=True)

            wb.save(output_path)
            wb.close()

            return True, f"تم تصدير {len(matching_orders)} طلب بنجاح إلى:\n{output_path}", len(matching_orders)

        except PermissionError:
            return False, "فشل التصدير: الملف المحدد مفتوح في برنامج آخر. يرجى إغلاقه أولاً.", 0
        except Exception as e:
            return False, f"خطأ أثناء تصدير الشيت: {str(e)}", 0
